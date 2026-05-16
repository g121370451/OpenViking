# -*- coding: utf-8 -*-
"""
生成 4 张分组柱状图：4 种实验模式 × 6 个数据集，对比 4 个指标维度。

数据源: ../relations-figure.xlsx (由 benchmark 实验汇总)
"""

import re
import math
import numpy as np
import matplotlib.pyplot as plt
import matplotlib
from openpyxl import load_workbook
from pathlib import Path

matplotlib.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei"]
matplotlib.rcParams["axes.unicode_minus"] = False
matplotlib.rcParams["font.family"] = "sans-serif"

EXCEL_PATH = Path(__file__).parent.parent / "relations-figure.xlsx"
OUTPUT_DIR = Path(__file__).parent / "output"
OUTPUT_DIR.mkdir(exist_ok=True)

MODES = [
    "agentic-learned",
    "agentic-unlearned",
    "unagentic-learned",
    "unagentic-unlearned",
]
DATASETS = ["FinanceBench", "QASPER", "SyllabusQA", "LocoMo", "NQ", "HotpotQA"]

# (列索引 0-based, 中文名, 英文文件名后缀, Y轴单位, 是否线性刻度)
METRICS = [
    (6, "Accuracy", "accuracy", "Accuracy", True),
    (7, "检索时间", "retrieval_time", "seconds/query", True),
    (8, "检索Token成本", "retrieval_cost", "tokens/query", True),
    (9, "迭代次数", "iterations", "iterations/query", True),
]

COLORS = ["#4E79A7", "#F28E2B", "#76B7B2", "#E15759"]


def parse_value(val):
    """从混合格式的单元格值中提取数值。"""
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val)
    s = re.sub(r"[\u200b-\u200f\u202a-\u202e\ufeff]", "", s)
    s = s.strip()
    if not s:
        return 0.0
    # 含等号: "68673+1934=70,607" → 取等号后的数值
    if "=" in s:
        after_eq = s.split("=")[-1]
        after_eq = re.sub(r"[^\d.]", "", after_eq)
        return float(after_eq) if after_eq else 0.0
    # 含单位: "67s", "0.45s", "62.02s"
    num_str = re.sub(r"[^\d.]", "", s)
    return float(num_str) if num_str else 0.0


def read_block_data(ws, start_row: int) -> dict[int, list[float]]:
    """读取一个区块的 6 行数据, 返回 {col_index: [6个数据集的数值]}."""
    result = {}
    for col in [6, 7, 8, 9]:  # accuracy, time, cost, iterations
        result[col] = []
    for r in range(start_row, start_row + 6):
        for col in [6, 7, 8, 9]:
            val = parse_value(ws.cell(row=r, column=col + 1).value)
            if col == 6:  # accuracy: 0.85 → 85
                val *= 100
            result[col].append(val)
    return result


def load_data():
    """从 Excel 读取四组实验数据, 返回 {mode: {col_idx: [6个值]}}."""
    wb = load_workbook(EXCEL_PATH, data_only=True)
    ws = wb["Sheet1"]

    # 四个区块的数据起始行 (1-indexed, 不含标题行)
    block_starts = {
        "agentic-learned": 2,       # 第 2-7 行
        "agentic-unlearned": 10,    # 第 10-15 行
        "unagentic-learned": 18,    # 第 18-23 行
        "unagentic-unlearned": 26,  # 第 26-31 行
    }

    data = {}
    for mode, start in block_starts.items():
        data[mode] = read_block_data(ws, start)

    wb.close()
    return data


def plot_metric(data, col_idx, title_cn, file_en, ylabel, linear_scale):
    """绘制单张柱状图。"""
    fig, ax = plt.subplots(figsize=(14, 7))

    n_modes = len(MODES)
    n_datasets = len(DATASETS)
    bar_width = 0.18
    x = np.arange(n_datasets)

    for i, mode in enumerate(MODES):
        values = data[mode][col_idx]
        offset = (i - n_modes / 2 + 0.5) * bar_width
        bars = ax.bar(
            x + offset, values, bar_width,
            label=mode, color=COLORS[i], edgecolor="white", linewidth=0.5,
        )
        # 柱顶标注数值
        for bar, val in zip(bars, values):
            if val > 0:
                if col_idx == 6:  # accuracy: 显示为百分比
                    label = f"{val:.0f}%"
                elif val < 100:
                    label = f"{val:.1f}"
                else:
                    label = f"{val:,.0f}"
                ax.text(
                    bar.get_x() + bar.get_width() / 2, bar.get_height(),
                    label, ha="center", va="bottom", fontsize=6,
                    rotation=45,
                )

    ax.set_title(f"Relations Experiment — {title_cn}", fontsize=16, fontweight="bold", pad=20)
    ax.set_ylabel(ylabel, fontsize=12)
    ax.set_xticks(x)
    ax.set_xticklabels(DATASETS, fontsize=11)

    if not linear_scale:
        ax.set_yscale("log")
        ymin = min(
            v for mode in MODES for v in data[mode][col_idx] if v > 0
        ) * 0.5
        ymax = max(
            v for mode in MODES for v in data[mode][col_idx]
        ) * 2
        ax.set_ylim(ymin, ymax)

    ax.grid(axis="y", linestyle="--", alpha=0.3)
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.08), ncol=4, frameon=False, fontsize=10)
    ax.set_axisbelow(True)

    plt.tight_layout()
    out_path = OUTPUT_DIR / f"chart_{file_en}.png"
    fig.savefig(out_path, dpi=150, bbox_inches="tight")
    plt.close(fig)
    print(f"  Saved: {out_path}")


def main():
    print("Loading data from:", EXCEL_PATH)
    data = load_data()

    # Print summary table
    print("\n=== Data Summary ===")
    for mode in MODES:
        print(f"\n{mode}:")
        for col, title, _, _, _ in METRICS:
            vals = [f"{v:.2f}" if v < 100 else f"{v:,.0f}" for v in data[mode][col]]
            print(f"  {title}: {vals}")

    print("\n=== Generating Charts ===")
    for col_idx, title_cn, file_en, ylabel, linear in METRICS:
        print(f"  Plotting: {title_cn}...")
        plot_metric(data, col_idx, title_cn, file_en, ylabel, linear)

    print(f"\nDone! Charts saved to: {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
